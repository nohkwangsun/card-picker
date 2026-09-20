#!/usr/bin/env python3
"""
카드 비교 xlsx -> data/cards_raw.json 변환기.

사용법:
    python3 scripts/xlsx_to_raw.py <입력.xlsx> [출력.json]

엑셀 구조 가정 (지금까지의 "카드 비교" 시트 형식):
  - "카드 비교" 라는 이름의 시트 (없으면 첫 번째 시트 사용)
  - A열에 '카드사'라는 라벨이 있는 행을 기준으로 카드 이름 행을 찾음
    (카드사 행 바로 위 행 = 카드 이름 행). 예전 레이아웃(1행=그룹, 2행=카드명,
    3행=카드사)이든 최신 레이아웃(1행=카드명, 2행=카드사, 그룹 행 없음)이든
    둘 다 이 규칙으로 자동 인식됨.
  - 카드 이름 행 바로 위에 병합된 행이 있으면 카드 그룹명(예: 프리미엄/일상/
    마일리지)으로 인식. 없으면 그룹 없이 처리.
  - A열: 각 행의 속성 이름 (카드사, 연회비, 가족카드, 연간 혜택, 혜택, 멤버십,
         마일리지, 할인, 해외 할인, 적립, 라운지, 공항, 발레파킹, Visa, 이벤트, 그외 등)

행/카드가 늘어나거나 순서가 바뀌거나 그룹 행이 통째로 없어져도, "카드사 라벨이
있는 행"을 기준점으로 삼아 카드 이름 행과 속성 행들을 다시 찾으므로 새 xlsx를
받으면 이 스크립트를 다시 돌리기만 하면 됩니다.
"""
import json
import sys
from pathlib import Path

import openpyxl

FIRST_CARD_COL = 2  # 카드 데이터가 시작하는 열 (B열)
LABEL_COL = 1      # 속성 이름(A열)


def find_sheet(wb):
    if "카드 비교" in wb.sheetnames:
        return wb["카드 비교"]
    for ws in wb.worksheets:
        if ws.max_row and ws.max_row > 1 and ws.max_column and ws.max_column > 1:
            return ws
    return wb.worksheets[0]


def find_name_row(ws):
    """A열에서 '카드사' 라벨이 있는 행을 찾아 그 바로 윗행(카드 이름 행)을 반환."""
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=LABEL_COL).value
        if v and str(v).strip() == "카드사" and r > 1:
            return r - 1
    raise ValueError("A열에서 '카드사' 라벨을 찾지 못했습니다 (카드 이름 행을 알 수 없음).")


def build_group_map(ws, name_row):
    """카드 이름 행 바로 위 행의 병합 범위를 이용해 열(column) -> 그룹명 매핑을 만든다.
    그룹 행이 없으면(name_row가 1행이거나 위 행에 병합/값이 없으면) 빈 매핑을 반환."""
    group_row = name_row - 1
    group_map = {}
    if group_row < 1:
        return group_map
    for mr in ws.merged_cells.ranges:
        if mr.min_row == group_row and mr.max_row == group_row:
            label = ws.cell(row=group_row, column=mr.min_col).value
            for c in range(mr.min_col, mr.max_col + 1):
                group_map[c] = label
    # 병합되지 않은 단일 셀 그룹 헤더도 지원
    for c in range(FIRST_CARD_COL, ws.max_column + 1):
        v = ws.cell(row=group_row, column=c).value
        if v and c not in group_map:
            group_map[c] = v
    return group_map


def collect_row_labels(ws, name_row):
    """A열에서 라벨이 있는 행 번호 -> 라벨명 (카드 이름 행 자체는 제외)."""
    labels = {}
    for r in range(1, ws.max_row + 1):
        if r == name_row:
            continue
        v = ws.cell(row=r, column=LABEL_COL).value
        if v:
            labels[r] = str(v).strip()
    return labels


def cell_value(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, str):
        v = v.strip()
        if v == "":
            return None
    return v


def extract(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = find_sheet(wb)
    name_row = find_name_row(ws)
    group_map = build_group_map(ws, name_row)
    row_labels = collect_row_labels(ws, name_row)

    cards = []
    for c in range(FIRST_CARD_COL, ws.max_column + 1):
        name = cell_value(ws, name_row, c)
        if not name:
            continue
        card = {
            "name": name,
            "group": group_map.get(c),
            "fields": {},
        }
        for r, label in sorted(row_labels.items()):
            v = cell_value(ws, r, c)
            if v is not None:
                card["fields"][label] = v
        cards.append(card)

    return {
        "sourceFile": Path(xlsx_path).name,
        "sheet": ws.title,
        "cardCount": len(cards),
        "cards": cards,
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else str(
        Path(__file__).resolve().parent.parent / "data" / "cards_raw.json"
    )
    data = extract(xlsx_path)
    Path(out_path).write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"OK: {data['cardCount']}개 카드 -> {out_path}")


if __name__ == "__main__":
    main()
